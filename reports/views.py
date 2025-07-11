from django.contrib.auth.mixins import LoginRequiredMixin
from django.views.generic import ListView, View
from orders.models import BOMFabric, BOMAccessory, Order, BOMVersion
from styles.models import Style, StyleCosting
from masters.models import Fabric, Accessory
from purchase.models import FabricPO, FabricReceiptItem, FabricReceipt, FabricPOItem
from django.db.models import Q, Sum
from django.http import HttpResponse
import csv
from collections import defaultdict
from django.shortcuts import render
from itertools import chain
from django.template.loader import get_template
from xhtml2pdf import pisa
from io import BytesIO
import openpyxl

class BOMReportView(LoginRequiredMixin, ListView):
    template_name = 'reports/bom_report.html'
    context_object_name = 'bom_entries'

    def get_queryset(self):
        style_id = self.request.GET.get('style')
        fabric_id = self.request.GET.get('fabric')
        order_id = self.request.GET.get('order')
        version_id = self.request.GET.get('version')

        fabric_bom = BOMFabric.objects.select_related('order', 'order__style', 'fabric', 'version').all()
        accessory_bom = BOMAccessory.objects.select_related('order', 'order__style', 'accessory', 'version').all()

        if style_id:
            fabric_bom = fabric_bom.filter(order__style__id=style_id)
            accessory_bom = accessory_bom.filter(order__style__id=style_id)
        if fabric_id:
            fabric_bom = fabric_bom.filter(fabric__id=fabric_id)
        if order_id:
            fabric_bom = fabric_bom.filter(order__id=order_id)
            accessory_bom = accessory_bom.filter(order__id=order_id)
        if version_id:
            fabric_bom = fabric_bom.filter(version__id=version_id)
            accessory_bom = accessory_bom.filter(version__id=version_id)

        # Combine and sort the querysets
        bom_entries = sorted(
            chain(fabric_bom, accessory_bom),
            key=lambda x: x.order.order_no
        )
        return bom_entries

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['all_styles'] = Style.objects.all()
        context['all_fabrics'] = Fabric.objects.all()
        context['all_orders'] = Order.objects.all()
        context['all_versions'] = BOMVersion.objects.all()
        context['selected_style'] = self.request.GET.get('style', '')
        context['selected_fabric'] = self.request.GET.get('fabric', '')
        context['selected_order'] = self.request.GET.get('order', '')
        context['selected_version'] = self.request.GET.get('version', '')
        context['report_format'] = self.request.GET.get('format', 'detailed')

        if context['report_format'] == 'summary':
            summary_data = defaultdict(lambda: {'required_qty': 0, 'issued_qty': 0, 'balance': 0})
            for entry in self.get_queryset():
                item_name = entry.fabric.name if isinstance(entry, BOMFabric) else entry.accessory.name
                item_type = 'Fabric' if isinstance(entry, BOMFabric) else 'Accessory'
                item_key = (item_type, item_name)
                summary_data[item_key]['required_qty'] += entry.required_qty
                summary_data[item_key]['issued_qty'] += entry.issued_qty
                summary_data[item_key]['balance'] += entry.balance
            context['summary_bom_entries'] = [{
                'type': k[0],
                'item_name': k[1],
                'required_qty': v['required_qty'],
                'issued_qty': v['issued_qty'],
                'balance': v['balance'],
            } for k, v in summary_data.items()]

        return context

    def render_to_pdf(self, template_src, context_dict={}):
        template = get_template(template_src)
        html = template.render(context_dict)
        result = BytesIO()
        pdf = pisa.pisaDocument(BytesIO(html.encode("UTF-8")), result)
        if not pdf.err:
            return HttpResponse(result.getvalue(), content_type='application/pdf')
        return None

    def get(self, request, *args, **kwargs):
        if 'export_pdf' in request.GET:
            queryset = self.get_queryset()
            context = {
                'bom_entries': queryset,
            }
            pdf = self.render_to_pdf('reports/bom_report_pdf.html', context)
            if pdf:
                response = HttpResponse(pdf, content_type='application/pdf')
                filename = "BOM_Report.pdf"
                content = "attachment; filename='%s'" %(filename)
                response['Content-Disposition'] = content
                return response
            return HttpResponse("Not found")

        if 'export_csv' in request.GET:
            response = HttpResponse(content_type='text/csv')
            response['Content-Disposition'] = 'attachment; filename="bom_report.csv"'

            writer = csv.writer(response)
            report_format = request.GET.get('format', 'detailed')

            if report_format == 'summary':
                writer.writerow(['Type', 'Item Name', 'Required Quantity', 'Issued Quantity', 'Balance'])
                context = self.get_context_data()
                for entry in context['summary_bom_entries']:
                    writer.writerow([
                        entry['type'],
                        entry['item_name'],
                        entry['required_qty'],
                        entry['issued_qty'],
                        entry['balance'],
                    ])
            else:
                writer.writerow(['Type', 'Order No', 'Style Name', 'Item Name', 'Required Quantity', 'Issued Quantity', 'Balance', 'BOM Version'])
                queryset = self.get_queryset()
                for entry in queryset:
                    item_name = entry.fabric.name if isinstance(entry, BOMFabric) else entry.accessory.name
                    item_type = 'Fabric' if isinstance(entry, BOMFabric) else 'Accessory'
                    writer.writerow([
                        item_type,
                        entry.order.order_no,
                        entry.order.style.name,
                        item_name,
                        entry.required_qty,
                        entry.issued_qty,
                        entry.balance,
                        entry.version.version_number if entry.version else '',
                    ])
            return response
        elif 'export_excel' in request.GET:
            response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
            response['Content-Disposition'] = 'attachment; filename="bom_report.xlsx"'

            workbook = openpyxl.Workbook()
            worksheet = workbook.active
            worksheet.title = "BOM Report"

            report_format = request.GET.get('format', 'detailed')

            if report_format == 'summary':
                worksheet.append(['Type', 'Item Name', 'Required Quantity', 'Issued Quantity', 'Balance'])
                context = self.get_context_data()
                for entry in context['summary_bom_entries']:
                    worksheet.append([
                        entry['type'],
                        entry['item_name'],
                        entry['required_qty'],
                        entry['issued_qty'],
                        entry['balance'],
                    ])
            else:
                worksheet.append(['Type', 'Order No', 'Style Name', 'Item Name', 'Required Quantity', 'Issued Quantity', 'Balance', 'BOM Version'])
                queryset = self.get_queryset()
                for entry in queryset:
                    item_name = entry.fabric.name if isinstance(entry, BOMFabric) else entry.accessory.name
                    item_type = 'Fabric' if isinstance(entry, BOMFabric) else 'Accessory'
                    worksheet.append([
                        item_type,
                        entry.order.order_no,
                        entry.order.style.name,
                        item_name,
                        entry.required_qty,
                        entry.issued_qty,
                        entry.balance,
                        entry.version.version_number if entry.version else '',
                    ])
            
            workbook.save(response)
            return response
        return super().get(request, *args, **kwargs)


class BOMDiffView(LoginRequiredMixin, View):
    template_name = 'reports/bom_diff.html'

    def get(self, request, *args, **kwargs):
        version1_id = request.GET.get('version1')
        version2_id = request.GET.get('version2')
        order_id = request.GET.get('order')

        version1 = None
        version2 = None
        order = None

        if order_id:
            order = Order.objects.filter(id=order_id).first()

        if version1_id:
            version1 = BOMVersion.objects.filter(id=version1_id).first()
        if version2_id:
            version2 = BOMVersion.objects.filter(id=version2_id).first()

        diff_results = []

        if version1 and version2 and order:
            # Get BOM items for version 1
            bom1_fabrics = {item.fabric.id: item for item in BOMFabric.objects.filter(order=order, version=version1)}
            bom1_accessories = {item.accessory.id: item for item in BOMAccessory.objects.filter(order=order, version=version1)}

            # Get BOM items for version 2
            bom2_fabrics = {item.fabric.id: item for item in BOMFabric.objects.filter(order=order, version=version2)}
            bom2_accessories = {item.accessory.id: item for item in BOMAccessory.objects.filter(order=order, version=version2)}

            # Compare fabrics
            all_fabric_ids = set(bom1_fabrics.keys()) | set(bom2_fabrics.keys())
            for fabric_id in all_fabric_ids:
                item1 = bom1_fabrics.get(fabric_id)
                item2 = bom2_fabrics.get(fabric_id)

                if item1 and item2:
                    if item1.required_qty != item2.required_qty:
                        diff_results.append({
                            'type': 'Fabric',
                            'item_name': item1.fabric.name,
                            'version1_qty': item1.required_qty,
                            'version2_qty': item2.required_qty,
                            'change': item2.required_qty - item1.required_qty
                        })
                elif item1:
                    diff_results.append({
                        'type': 'Fabric',
                        'item_name': item1.fabric.name,
                        'version1_qty': item1.required_qty,
                        'version2_qty': 0,
                        'change': -item1.required_qty,
                        'status': 'Removed'
                    })
                elif item2:
                    diff_results.append({
                        'type': 'Fabric',
                        'item_name': item2.fabric.name,
                        'version1_qty': 0,
                        'version2_qty': item2.required_qty,
                        'change': item2.required_qty,
                        'status': 'Added'
                    })

            # Compare accessories
            all_accessory_ids = set(bom1_accessories.keys()) | set(bom2_accessories.keys())
            for accessory_id in all_accessory_ids:
                item1 = bom1_accessories.get(accessory_id)
                item2 = bom2_accessories.get(accessory_id)

                if item1 and item2:
                    if item1.required_qty != item2.required_qty:
                        diff_results.append({
                            'type': 'Accessory',
                            'item_name': item1.accessory.name,
                            'version1_qty': item1.required_qty,
                            'version2_qty': item2.required_qty,
                            'change': item2.required_qty - item1.required_qty
                        })
                elif item1:
                    diff_results.append({
                        'type': 'Accessory',
                        'item_name': item1.accessory.name,
                        'version1_qty': item1.required_qty,
                        'version2_qty': 0,
                        'change': -item1.required_qty,
                        'status': 'Removed'
                    })
                elif item2:
                    diff_results.append({
                        'type': 'Accessory',
                        'item_name': item2.accessory.name,
                        'version1_qty': 0,
                        'version2_qty': item2.required_qty,
                        'change': item2.required_qty,
                        'status': 'Added'
                    })

        context = {
            'all_orders': Order.objects.all(),
            'selected_order': order_id,
            'all_versions': BOMVersion.objects.filter(order=order).order_by('-version_number') if order else BOMVersion.objects.none(),
            'selected_version1': version1_id,
            'selected_version2': version2_id,
            'diff_results': diff_results,
            'version1_obj': version1,
            'version2_obj': version2,
        }
        return render(request, self.template_name, context)

class PurchaseOrderReportView(LoginRequiredMixin, ListView):
    template_name = 'reports/purchase_order_report.html'
    context_object_name = 'purchase_orders'

    def get_queryset(self):
        queryset = FabricPO.objects.all()
        status = self.request.GET.get('status')
        start_date = self.request.GET.get('start_date')
        end_date = self.request.GET.get('end_date')
        po_id = self.request.GET.get('po_id')

        if po_id:
            queryset = queryset.filter(id=po_id)
        if status:
            queryset = queryset.filter(status=status)
        if start_date:
            queryset = queryset.filter(date__gte=start_date)
        if end_date:
            queryset = queryset.filter(date__lte=end_date)

        for po in queryset:
            po.total_received_qty = FabricReceiptItem.objects.filter(receipt__po_ref=po).aggregate(Sum('quantity_received'))['quantity_received__sum'] or 0
            po.pending_qty = po.total_qty - po.total_received_qty
        return queryset

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['all_statuses'] = FabricPO.STATUS_CHOICES
        context['selected_status'] = self.request.GET.get('status', '')
        context['start_date'] = self.request.GET.get('start_date', '')
        context['end_date'] = self.request.GET.get('end_date', '')
        return context

    def get(self, request, *args, **kwargs):
        if 'print_po' in request.GET:
            po_id = request.GET.get('po_id')
            if po_id:
                po = self.get_queryset().first()
                if po:
                    return render(request, 'reports/purchase_order_print.html', {'po': po})
            # Handle case where po_id is not provided or PO not found
            return super().get(request, *args, **kwargs)
        return super().get(request, *args, **kwargs)

class PurchaseReceiptReportView(LoginRequiredMixin, ListView):
    template_name = 'reports/purchase_receipt_report.html'
    context_object_name = 'purchase_receipts'

    def get_queryset(self):
        queryset = FabricReceipt.objects.all()
        start_date = self.request.GET.get('start_date')
        end_date = self.request.GET.get('end_date')

        if start_date:
            queryset = queryset.filter(receipt_date__gte=start_date)
        if end_date:
            queryset = queryset.filter(receipt_date__lte=end_date)

        for receipt in queryset:
            receipt.total_ordered_qty = FabricPOItem.objects.filter(po=receipt.po_ref).aggregate(Sum('quantity'))['quantity__sum'] or 0
            receipt.total_received_qty = FabricReceiptItem.objects.filter(receipt=receipt).aggregate(Sum('quantity_received'))['quantity_received__sum'] or 0
            receipt.is_delayed = receipt.receipt_date > receipt.po_ref.delivery_date
            receipt.is_short = receipt.total_received_qty < receipt.total_ordered_qty
        return queryset

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['start_date'] = self.request.GET.get('start_date', '')
        context['end_date'] = self.request.GET.get('end_date', '')
        context['report_format'] = self.request.GET.get('format', 'detailed')

        if context['report_format'] == 'summary':
            summary_data = defaultdict(lambda: {'ordered_qty': 0, 'received_qty': 0})
            for receipt in self.get_queryset(): # Use get_queryset to get filtered data
                for item in receipt.items.all():
                    item_key = (item.fabric.supplier, item.fabric.name)
                    summary_data[item_key]['ordered_qty'] += item.quantity_received # This should be ordered from POItem
                    summary_data[item_key]['received_qty'] += item.quantity_received
            context['summary_receipt_entries'] = [{
                'supplier': k[0],
                'fabric_name': k[1],
                'ordered_qty': v['ordered_qty'],
                'received_qty': v['received_qty'],
            } for k, v in summary_data.items()]

        return context

    def get(self, request, *args, **kwargs):
        if 'export_csv' in request.GET:
            response = HttpResponse(content_type='text/csv')
            response['Content-Disposition'] = 'attachment; filename="purchase_receipt_report.csv"'

            writer = csv.writer(response)
            report_format = request.GET.get('format', 'detailed')

            if report_format == 'summary':
                writer.writerow(['Supplier', 'Fabric Name', 'Ordered Quantity', 'Received Quantity'])
                context = self.get_context_data()
                for entry in context['summary_receipt_entries']:
                    writer.writerow([
                        entry['supplier'],
                        entry['fabric_name'],
                        entry['ordered_qty'],
                        entry['received_qty'],
                    ])
            else:
                writer.writerow(['GRN No', 'PO Reference', 'Receipt Date', 'Ordered Quantity', 'Received Quantity', 'Status'])
                queryset = self.get_queryset()
                for receipt in queryset:
                    writer.writerow([
                        receipt.grn_no,
                        receipt.po_ref.po_no,
                        receipt.receipt_date,
                        receipt.total_ordered_qty,
                        receipt.total_received_qty,
                        ('Delayed ' if receipt.is_delayed else '') + ('Short ' if receipt.is_short else '') + ('On Time & Full' if not receipt.is_delayed and not receipt.is_short else ''),
                    ])
            return response
        elif 'export_excel' in request.GET:
            response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
            response['Content-Disposition'] = 'attachment; filename="purchase_receipt_report.xlsx"'

            workbook = openpyxl.Workbook()
            worksheet = workbook.active
            worksheet.title = "Purchase Receipt Report"

            report_format = request.GET.get('format', 'detailed')

            if report_format == 'summary':
                worksheet.append(['Supplier', 'Fabric Name', 'Ordered Quantity', 'Received Quantity'])
                context = self.get_context_data()
                for entry in context['summary_receipt_entries']:
                    worksheet.append([
                        entry['supplier'],
                        entry['fabric_name'],
                        entry['ordered_qty'],
                        entry['received_qty'],
                    ])
            else:
                worksheet.append(['GRN No', 'PO Reference', 'Receipt Date', 'Ordered Quantity', 'Received Quantity', 'Status'])
                queryset = self.get_queryset()
                for receipt in queryset:
                    worksheet.append([
                        receipt.grn_no,
                        receipt.po_ref.po_no,
                        receipt.receipt_date,
                        receipt.total_ordered_qty,
                        receipt.total_received_qty,
                        ('Delayed ' if receipt.is_delayed else '') + ('Short ' if receipt.is_short else '') + ('On Time & Full' if not receipt.is_delayed and not receipt.is_short else ''),
                    ])
            
            workbook.save(response)
            return response
        elif 'print_receipt' in request.GET:
            receipt_id = request.GET.get('receipt_id')
            if receipt_id:
                receipt = self.get_queryset().filter(id=receipt_id).first()
                if receipt:
                    return render(request, 'reports/purchase_receipt_print.html', {'receipt': receipt})
            return super().get(request, *args, **kwargs)
        return super().get(request, *args, **kwargs)

class StyleCostingReportView(LoginRequiredMixin, ListView):
    model = StyleCosting
    template_name = 'reports/style_costing_report.html'
    context_object_name = 'style_costings'

    def render_to_pdf(self, template_src, context_dict={}):
        template = get_template(template_src)
        html = template.render(context_dict)
        result = BytesIO()
        pdf = pisa.pisaDocument(BytesIO(html.encode("UTF-8")), result)
        if not pdf.err:
            return HttpResponse(result.getvalue(), content_type='application/pdf')
        return None

    def get_queryset(self):
        queryset = super().get_queryset()
        style_id = self.request.GET.get('style')
        if style_id:
            queryset = queryset.filter(style__id=style_id)
        return queryset

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['all_styles'] = Style.objects.all()
        context['selected_style'] = self.request.GET.get('style', '')
        return context

    def get(self, request, *args, **kwargs):
        if 'export_csv' in request.GET:
            response = HttpResponse(content_type='text/csv')
            response['Content-Disposition'] = 'attachment; filename="style_costing_report.csv"'

            writer = csv.writer(response)
            writer.writerow([
                'Style Name', 'Currency', 'Exchange Rate', 'Total Fabric Cost',
                'Total Accessory Cost', 'Total Cost', 'Margin/Markup (%)',
                'Discount (%)', 'Landed Cost', 'Retail Price'
            ])

            queryset = self.get_queryset()
            for costing in queryset:
                writer.writerow([
                    costing.style.name,
                    costing.currency.code if costing.currency else '',
                    costing.exchange_rate,
                    costing.total_fabric_cost,
                    costing.total_accessory_cost,
                    costing.total_cost,
                    costing.margin_markup,
                    costing.discount,
                    costing.landed_cost,
                    costing.retail_price,
                ])
            return response
        if 'export_pdf' in request.GET:
            queryset = self.get_queryset()
            context = {
                'style_costings': queryset,
                'all_styles': Style.objects.all(), # Pass all_styles for filter display in PDF if needed
                'selected_style': self.request.GET.get('style', ''),
            }
            pdf = self.render_to_pdf('reports/style_costing_report_pdf.html', context)
            if pdf:
                response = HttpResponse(pdf, content_type='application/pdf')
                filename = "Style_Costing_Report.pdf"
                content = "attachment; filename='%s'" %(filename)
                response['Content-Disposition'] = content
                return response
            return HttpResponse("Not found")
        return super().get(request, *args, **kwargs)